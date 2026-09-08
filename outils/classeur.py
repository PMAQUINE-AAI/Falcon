"""Construit le classeur FALCON. Outil de CONSTRUCTION, pas un livrable.

    python outils/classeur.py

**`openpyxl` n'entre pas dans le produit.** Il s'installe ici pour fabriquer
le fichier, et il n'est ni dans `pyproject.toml`, ni dans le bundle. La regle
« PyYAML seule dependance d'execution » tient : FALCON ne lit jamais un
classeur, il lit les trois CSV que le classeur ecrit.

**Pourquoi un `.xlsx` et pas un `.xlsm`.** Un projet VBA est un flux binaire
OLE (`vbaProject.bin`) ; `openpyxl` sait le CONSERVER dans un fichier qui en a
deja un, il ne sait pas en CREER. Fabriquer ce binaire a la main serait
possible en theorie et invérifiable en pratique : je ne peux ouvrir Excel
nulle part ici, donc je ne pourrais pas etablir que le fichier s'ouvre. Livrer
un binaire dont on ne sait pas s'il s'ouvre est exactement le defaut que ce
depot traque — un artefact d'aspect normal, faux, et dont l'erreur ne se
verrait qu'au moment ou quelqu'un en a besoin.

Le classeur est donc livre en `.xlsx`, avec les modules `.bas` a importer en
deux gestes. Voir `classeur/LISEZMOI.md`.

**La parade au VBA non teste est architecturale : le garder bete.** La macro
ecrit trois CSV, sans rien valider. Toute la validation vit dans
`falcon/tableur/`, qui est teste. Une macro fausse produit donc un CSV refuse
avec un message situe — pas un YAML plausible et faux.
"""

from __future__ import annotations

import sys
from pathlib import Path

RACINE = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RACINE))

from falcon.noyau import COMPARAISONS, DEROGEABLES                # noqa: E402
from falcon.pipeline.composition import NOMS as TRANSFORMATIONS   # noqa: E402
from falcon.pipeline.modele import ACTIONS, GENRES_SOURCE         # noqa: E402
from falcon.tableur.lecture import (                              # noqa: E402
    COLONNES_DEROGATIONS, COLONNES_ETAPES, COLONNES_PIPELINE, JETON_LIBRE,
)

SORTIE = RACINE / "classeur" / "FALCON.xlsx"

#: Colonnes du dictionnaire exporte par `falcon dictionnaire`. On les recopie
#: plutot que d'importer `commandes.dictionnaire` : ce module-la vit sous
#: `falcon/commandes/`, et l'outil de construction n'a pas a en dependre.
COLONNES_DICTIONNAIRE = (
    "ecran", "cible", "transaction", "programme", "dynpro", "empreinte",
    "titre", "type", "soustype", "nom", "texte", "modifiable", "infobulle",
    "releve",
)

#: Les colonnes dont la valeur part dans SAP, et qui s'ecrivent donc entre
#: crochets. Elles sont mises au format TEXTE en plus — ceinture et bretelles,
#: parce que le format se perd au copier-coller et pas les crochets.
ENCADREES = ("source_valeur", "defaut", "statut_attendu")


def _feuille_libellee(classeur, titre: str, colonnes, aide: str):
    """Une feuille a en-tete fige, avec sa ligne d'explication."""
    from openpyxl.styles import Alignment, Font, PatternFill

    feuille = classeur.create_sheet(titre)
    feuille["A1"] = aide
    feuille["A1"].font = Font(italic=True, size=9)
    feuille["A1"].alignment = Alignment(vertical="top")

    for rang, nom in enumerate(colonnes, start=1):
        cellule = feuille.cell(row=2, column=rang, value=nom)
        cellule.font = Font(bold=True)
        cellule.fill = PatternFill("solid", fgColor="DDE7F0")
        feuille.column_dimensions[cellule.column_letter].width = max(
            12, min(34, len(nom) + 8))

    # L'en-tete est en ligne 2 (la 1 porte l'aide) et le volet se fige sous
    # elle : sur trois cents lignes d'etapes, perdre les noms de colonne fait
    # remplir la mauvaise.
    feuille.freeze_panes = "A3"
    return feuille


def _liste(feuille, colonne: str, valeurs, message: str,
           premier: int = 3, dernier: int = 500) -> None:
    """Une liste deroulante sur une colonne entiere.

    `allow_blank` reste vrai : beaucoup de colonnes sont facultatives, et une
    validation qui refuse le vide empecherait de saisir une etape simple.
    """
    from openpyxl.worksheet.datavalidation import DataValidation

    # `showErrorMessage` et `showInputMessage` doivent etre POSES.
    #
    # Ils valent 0 par defaut, et `openpyxl` ne les leve pas quand on
    # renseigne `error`/`prompt` : le XML sortait avec
    # `showErrorMessage="0" showInputMessage="0"`, donc l'aide ne s'affichait
    # jamais et une valeur hors liste etait ACCEPTEE sans alerte. Seule la
    # fleche de la liste subsistait.
    #
    # C'est la seule documentation qui vive DANS la feuille — « constante :
    # ecrite ici, ENTRE CROCHETS », « pas VRAI/FAUX », « zeros:18 » — et elle
    # etait invisible. La justification de `allow_blank` ci-dessus decrivait
    # meme un reglage dont l'effet n'existait pas.
    validation = DataValidation(
        type="list", formula1='"' + ",".join(valeurs) + '"',
        allow_blank=True, showDropDown=False,
        showErrorMessage=True, showInputMessage=True)
    validation.error = message
    validation.errorTitle = "Valeur hors liste"
    validation.prompt = message
    validation.promptTitle = "Choisir"
    feuille.add_data_validation(validation)
    validation.add(f"{colonne}{premier}:{colonne}{dernier}")


def _texte(feuille, colonnes: dict[str, str],
           premier: int = 3, dernier: int = 500) -> None:
    """Format TEXTE sur les colonnes dont Excel retyperait la valeur."""
    for lettre in colonnes.values():
        for rang in range(premier, dernier + 1):
            feuille[f"{lettre}{rang}"].number_format = "@"


def construire(sortie: Path = SORTIE) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    classeur = Workbook()
    classeur.remove(classeur.active)

    # -- 1. Dictionnaire : ce qu'on COLLE depuis `falcon dictionnaire` ------
    dictionnaire = _feuille_libellee(
        classeur, "Dictionnaire", COLONNES_DICTIONNAIRE,
        "COLLE ICI le CSV produit par « falcon dictionnaire ». Les colonnes "
        "`ecran` et `cible` sont celles qu'on recopie dans Etapes — on les "
        "colle, on ne les retape pas.")

    # -- 2. Pipeline : clef / valeur ---------------------------------------
    pipeline = _feuille_libellee(
        classeur, "Pipeline", COLONNES_PIPELINE,
        "Les proprietes de la pipeline. `cles` accepte plusieurs colonnes "
        "separees par des virgules.")
    for rang, (nom, indice) in enumerate((
            ("nom", "sans espace ni accent, c'est lui qu'on tapera pour "
                    "confirmer"),
            ("classe", "iterative"),
            ("cles", "les colonnes du jeu qui identifient un item"),
            ("plafond_items", "OBLIGATOIRE — le rayon d'action du lot"),
            ("plafond_sauvegardes", "OBLIGATOIRE — ce qui part dans SAP"),
    ), start=3):
        pipeline[f"A{rang}"] = nom
        pipeline[f"C{rang}"] = indice
        pipeline[f"C{rang}"].font = Font(italic=True, size=9, color="808080")
    pipeline.column_dimensions["C"].width = 52
    for rang in range(3, 8):
        pipeline[f"B{rang}"].number_format = "@"

    # -- 3. Etapes ---------------------------------------------------------
    etapes = _feuille_libellee(
        classeur, "Etapes", COLONNES_ETAPES,
        "Une ligne par etape. Le `rang` doit CROITRE (numerote de 10 en 10 "
        "pour pouvoir inserer) : trier sur une autre colonne reordonnerait "
        "les etapes en silence. Les valeurs tapees dans SAP s'ecrivent ENTRE "
        "CROCHETS — [0100], [S], [] pour la chaine vide.")
    lettres = {nom: etapes.cell(row=2, column=rang).column_letter
               for rang, nom in enumerate(COLONNES_ETAPES, start=1)}

    _liste(etapes, lettres["action"], sorted(ACTIONS),
           "Ce que l'etape fait. `set` ecrit, `press` clique, `vkey` envoie "
           "une touche de fonction, `lire` releve une valeur.")
    _liste(etapes, lettres["source_genre"], sorted(GENRES_SOURCE),
           "D'ou vient la valeur. `colonne` : du jeu. `constante` : ecrite "
           "ici, ENTRE CROCHETS. `gabarit` : composee, « /BCP01_{site} ». "
           "`lue` : le nom d'une etape `lire` precedente.")
    _liste(etapes, lettres["sauvegarde"], ("oui", "non"),
           "« oui » si cette etape VALIDE dans SAP. Pas VRAI/FAUX : le "
           "booleen d'Excel est localise et ne se relit pas d'une machine a "
           "l'autre.")
    _liste(etapes, lettres["comparaison"], sorted(COMPARAISONS),
           "Comment la garde de relecture compare ce qui a ete tape a ce que "
           "SAP rend. Vide = `casse`.")
    _liste(etapes, lettres["format"], sorted(TRANSFORMATIONS),
           "Transformations appliquees DANS L'ORDRE, separees par des "
           "virgules. « zeros » et « tronque » prennent une largeur : "
           "« zeros:18 ».")
    _texte(etapes, {c: lettres[c] for c in ENCADREES})
    # `ecran` et `cible` se collent depuis le Dictionnaire : format texte,
    # pour qu'un jeton « IA08::RIPLKO10::0100 » ne soit pas reinterprete.
    _texte(etapes, {c: lettres[c] for c in ("ecran", "cible")})

    # L'exemple vit sur SA feuille, pas dans `Etapes`.
    #
    # Une ligne d'exemple laissee dans la feuille exportee deviendrait une
    # ETAPE : la macro l'ecrirait dans le CSV, et la pipeline porterait une
    # saisie que personne n'a voulue. C'est le genre de piege qu'on ne voit
    # qu'une fois — apres l'avoir joue.
    modele_ = classeur.create_sheet("Exemple")
    modele_["A1"] = ("CETTE FEUILLE N'EST PAS EXPORTEE. Copie les lignes vers "
                     "« Etapes » et adapte-les.")
    modele_["A1"].font = Font(italic=True, size=9)
    for rang, nom in enumerate(COLONNES_ETAPES, start=1):
        cellule = modele_.cell(row=2, column=rang, value=nom)
        cellule.font = Font(bold=True)
        modele_.column_dimensions[cellule.column_letter].width = max(
            12, min(34, len(nom) + 8))
    for decalage, ligne in enumerate((
            ("10", "saisir_site", "set", "wnd[0]/usr/ctxtWERKS-LOW",
             "IA08::RIPLKO10::1000", "colonne", "site", "majuscules", "",
             "", "non", "", "", ""),
            ("20", "saisir_variante", "set", "wnd[0]/usr/ctxtVARIANT",
             "IA08::RIPLKO10::1000", "gabarit", "/BCP01_{site}", "", "",
             "", "non", "", "", ""),
            ("30", "saisir_equipement", "set", "wnd[0]/usr/ctxtEQUNR",
             "IA08::RIPLKO10::1000", "colonne", "equipement",
             "sans_espaces_autour, zeros:18", "", "", "non", "", "", ""),
            ("40", "sauver", "press", "wnd[0]/tbar[0]/btn[11]",
             "IA08::RIPLKO10::1000", "", "", "", "", "[S]", "oui", "", "", ""),
    )):
        for colonne, valeur in enumerate(ligne, start=1):
            modele_.cell(row=3 + decalage, column=colonne, value=valeur)
    # Le « # » de tete n'est pas decoratif : c'est la marque de commentaire
    # que la macro saute. Une note libre sous un tableau exporte deviendrait
    # une ETAPE dont le `rang` serait une phrase.
    modele_[f"A{3 + 4 + 1}"] = (
        f"# Dans `ecran` : le jeton colle depuis Dictionnaire, ou "
        f"« {JETON_LIBRE} » si l'etape ne sait pas encore ou elle atterrit.")
    modele_[f"A{3 + 4 + 1}"].font = Font(italic=True, size=9, color="808080")

    # -- 4. Derogations ----------------------------------------------------
    derogations = _feuille_libellee(
        classeur, "Derogations", COLONNES_DEROGATIONS,
        "Une ligne par derogation. `etape` doit nommer une etape EXISTANTE — "
        "une derogation orpheline ne couvrirait rien, et la garde resterait "
        "armee. Le motif fait 30 caracteres minimum : une derogation non "
        "expliquee est une garde desactivee en douce.")
    _liste(derogations, "B", sorted(DEROGEABLES),
           "Seules ces trois gardes se derogent. Une identite violee signale "
           "que le modele du monde est faux, et le plafond est la derniere "
           "barriere.")
    derogations.column_dimensions["D"].width = 70

    # -- 5. Donnees : le jeu -----------------------------------------------
    donnees = classeur.create_sheet("Donnees")
    donnees["A1"] = ("Le jeu. UNE LIGNE = UNE ITERATION. Les en-tetes sont "
                     "les noms de colonne que les etapes citent.")
    donnees["A1"].font = Font(italic=True, size=9)
    donnees["A2"] = "site"
    donnees["B2"] = "equipement"
    for lettre in ("A", "B", "C", "D", "E", "F", "G", "H"):
        donnees.column_dimensions[lettre].width = 18
        for rang in range(3, 1001):
            donnees[f"{lettre}{rang}"].number_format = "@"
    donnees.freeze_panes = "A3"

    sortie.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(sortie)
    return sortie


def main() -> int:
    chemin = construire()
    taille = chemin.stat().st_size
    print(f"{chemin}  ({taille // 1024} Kio)")
    print("\n  C'est un .xlsx : les macros s'importent depuis classeur/*.bas.")
    print("  Voir classeur/LISEZMOI.md — et la raison y est dite en toutes")
    print("  lettres, plutot que decouverte a l'ouverture.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
