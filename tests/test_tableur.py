"""Du classeur a la pipeline. Ce qui est refuse compte plus que ce qui passe.

Ce convertisseur est le chainon qui rend vraie la phrase « un automatisme
nouveau est un fichier, pas un commit ». Il est donc, aussi, le nouvel endroit
ou une valeur peut se faire retyper en silence entre le tableur et SAP — et
c'est ce que ces tests surveillent en premier.
"""

from __future__ import annotations

import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from falcon.pipeline import charger
from falcon.tableur import (
    COLONNES_DEROGATIONS, COLONNES_ETAPES, CSV_DEROGATIONS, CSV_ETAPES,
    CSV_PIPELINE, TableurInvalide, convertir, convertir_fichiers,
)

PIPELINE = """propriete;valeur
nom;bcp_ia08_variantes
classe;iterative
cles;site
plafond_items;50
plafond_sauvegardes;50
"""

ENTETE = ("rang;nom;action;cible;ecran;source_genre;source_valeur;format;"
          "defaut;statut_attendu;sauvegarde;comparaison;fenetres;fonction\n")

ECRAN = "IA08::RIPLKO10::1000"

ETAPES = ENTETE + (
    f"10;saisir;set;wnd[0]/usr/ctxtWERKS-LOW;{ECRAN};colonne;site;;;;non;;;\n"
    f"20;sauver;press;wnd[0]/tbar[0]/btn[11];{ECRAN};;;;;;oui;;;\n"
)


class Base(unittest.TestCase):

    def setUp(self):
        dossier = tempfile.TemporaryDirectory()
        self.addCleanup(dossier.cleanup)
        self.racine = Path(dossier.name)
        self._poser(CSV_PIPELINE, PIPELINE)
        self._poser(CSV_ETAPES, ETAPES)

    def _poser(self, nom: str, contenu: str) -> Path:
        chemin = self.racine / nom
        chemin.write_text(contenu, encoding="utf-8")
        return chemin

    def _etapes(self, *lignes: str) -> None:
        self._poser(CSV_ETAPES, ENTETE + "".join(l + "\n" for l in lignes))


class TestConversionNominale(Base):

    def test_les_trois_CSV_donnent_une_pipeline_qui_CHARGE(self):
        """Le seul critere qui compte : ce qui sort doit passer `charger()`."""
        chemin = convertir_fichiers(self.racine, self.racine / "p.yaml")
        pipeline = charger(chemin)
        self.assertEqual(pipeline.nom, "bcp_ia08_variantes")
        self.assertEqual(pipeline.cles, ("site",))
        self.assertEqual([e.nom for e in pipeline.etapes],
                         ["saisir", "sauver"])
        self.assertTrue(pipeline.etapes[1].sauvegarde)

    def test_l_ecran_tient_dans_UNE_cellule(self):
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        self.assertEqual(pipeline.etapes[0].ecran,
                         ("IA08", "RIPLKO10", "1000"))

    def test_le_jeton_libre_est_l_AUTRE_branche_du_XOR(self):
        """Le XOR `ecran` / `navigation_libre` devient INEXPRIMABLE plutot que
        valide : une cellule porte un triplet, ou le mot « libre ». On ne peut
        pas declarer les deux, donc on ne peut pas se les voir refuser."""
        self._etapes("10;aller;press;wnd[0]/tbar[0]/btn[3];libre;;;;;;non;;;")
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        self.assertIsNone(pipeline.etapes[0].ecran)
        self.assertTrue(pipeline.etapes[0].navigation_libre)

    def test_un_gabarit_traverse_la_conversion(self):
        """Le cas qui a motive `composition` : « /BCP01_{site} »."""
        self._etapes(
            f"10;variante;set;wnd[0]/usr/ctxtV;{ECRAN};gabarit;"
            f"/BCP01_{{site}};;;;non;;;")
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        source = pipeline.etapes[0].source
        self.assertEqual((source.genre, source.valeur),
                         ("gabarit", "/BCP01_{site}"))
        self.assertEqual(source.colonnes, ("site",))

    def test_le_format_traverse_dans_l_ORDRE_declare(self):
        self._etapes(
            f"10;num;set;wnd[0]/usr/ctxtM;{ECRAN};colonne;site;"
            f"majuscules, zeros:12;;;non;;;")
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        self.assertEqual([str(f) for f in pipeline.etapes[0].format],
                         ["majuscules", "zeros: 12"])

    def test_une_derogation_se_rattache_a_son_etape(self):
        self._poser(CSV_DEROGATIONS,
                    "etape;garde;portee;motif\n"
                    "sauver;relecture;;L'ecran de confirmation n'expose pas "
                    "le champ relu, releve le 2026-09-08.\n")
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        derogations = pipeline.etapes[1].derogations
        self.assertEqual(len(derogations), 1)
        self.assertEqual(derogations[0].garde, "relecture")
        # La portee vide se remplit avec l'etape porteuse : sans ca, elle
        # vaudrait « * » et relacherait la garde sur TOUTE la pipeline.
        self.assertEqual(derogations[0].portee, "etape:sauver")

    def test_le_BOM_d_Excel_ne_fabrique_pas_une_colonne_fantome(self):
        """`falcon dictionnaire` ecrit un BOM, et Excel le reecrit. Laisse
        dans la premiere cellule, il ferait une colonne « ﻿rang »."""
        (self.racine / CSV_ETAPES).write_bytes(
            b"\xef\xbb\xbf" + ETAPES.encode("utf-8"))
        self.assertIn("saisir", convertir(self.racine))

    def test_la_virgule_marche_aussi_comme_delimiteur(self):
        self._poser(CSV_PIPELINE, PIPELINE.replace(";", ","))
        self._poser(CSV_ETAPES, ETAPES.replace(";", ","))
        self.assertIn("bcp_ia08_variantes", convertir(self.racine))


class TestDeterminisme(Base):
    """L'empreinte est un `sha256` du TEXTE INTEGRAL, et la reprise en depend.

    Deux conversions du meme CSV qui differeraient d'une espace donneraient
    deux empreintes, donc une reprise refusee sur une pipeline pourtant
    identique — et l'utilisateur apprendrait a passer outre avec `forcer`, ce
    qui desarmerait la garde pour de bon.
    """

    def test_deux_conversions_donnent_les_MEMES_octets(self):
        self.assertEqual(convertir(self.racine), convertir(self.racine))

    def test_l_empreinte_ne_bouge_pas_d_une_conversion_a_l_autre(self):
        premiere = charger(convertir_fichiers(self.racine,
                                              self.racine / "a.yaml"))
        seconde = charger(convertir_fichiers(self.racine,
                                             self.racine / "b.yaml"))
        self.assertEqual(premiere.empreinte, seconde.empreinte)

    def test_rien_du_MONDE_n_entre_dans_la_sortie(self):
        """Ni horodatage, ni chemin source, ni version : trois choses qu'il
        serait naturel de mettre en en-tete et qui feraient changer
        l'empreinte a chaque conversion."""
        texte = convertir(self.racine)
        self.assertNotIn(str(self.racine), texte)
        self.assertFalse(texte.startswith("#"),
                         "un en-tete de commentaire entrerait dans l'empreinte")


class TestLesCrochets(Base):
    """La convention qui empeche Excel de retyper une valeur destinee a SAP.

    Prescrire un format « Texte » ne suffirait pas : FALCON ne voit jamais le
    classeur, et un « 0100 » deja mange par Excel est indiscernable d'un
    « 100 » legitime — il n'y a plus rien a rattraper a la lecture.
    """

    def test_une_constante_SANS_crochets_est_refusee(self):
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};constante;0100;;;;non;;;")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        message = str(capture.exception)
        self.assertIn("crochets", message)
        self.assertIn("[0100]", message)

    def test_une_constante_AVEC_crochets_garde_ses_zeros(self):
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};constante;[0100];;;;non;;;")
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        self.assertEqual(pipeline.etapes[0].source.valeur, "0100")

    def test_les_crochets_vides_expriment_la_chaine_VIDE(self):
        """« [] » est distinct de « pas de defaut » : c'est la difference
        entre « vide le champ » et « ne fais rien de particulier »."""
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};colonne;site;;[];;non;;;")
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        self.assertEqual(pipeline.etapes[0].defaut, "")

    def test_un_NOM_DE_COLONNE_ne_prend_PAS_de_crochets(self):
        """Une reference interne n'est jamais tapee dans SAP. L'encadrer ferait
        chercher une colonne nommee « [site] »."""
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};colonne;[site];;;;non;;;")
        pipeline = charger(convertir_fichiers(self.racine,
                                              self.racine / "p.yaml"))
        self.assertEqual(pipeline.etapes[0].source.valeur, "[site]")

    def test_un_caractere_qui_est_un_SAUT_DE_LIGNE_pour_YAML_survit(self):
        """U+0085 (NEL) n'etait pas echappe, et PyYAML le replie en ESPACE a
        l'interieur d'un scalaire entre guillemets.

        Mesure avant correction, sur une constante destinee a SAP :

            cellule  'K75\x85B'   ->   tape dans SAP  'K75 B'

        La valeur entre crochets — la seule forme censee survivre a tout —
        etait donc tapee alteree. Et la garde de relecture ne voyait rien :
        elle compare la valeur TRANSFORMEE a ce que SAP rend.
        """
        for nom, caractere in (("U+0085 NEL", "\u0085"),
                               ("U+2028 LS", "\u2028"),
                               ("U+2029 PS", "\u2029"),
                               ("U+00A0 NBSP", "\u00a0")):
            with self.subTest(caractere=nom):
                valeur = "K75" + caractere + "B"
                self._etapes(
                    f"10;a;set;wnd[0]/usr/c;{ECRAN};constante;[{valeur}];"
                    f";;;non;;;",
                    f"99;sauver;press;wnd[0]/tbar[0]/btn[11];{ECRAN};;;;;;"
                    f"oui;;;")
                pipeline = charger(convertir_fichiers(self.racine,
                                                      self.racine / "p.yaml"))
                self.assertEqual(pipeline.etapes[0].source.valeur, valeur,
                                 f"{nom} altere entre la cellule et SAP")

    def test_un_statut_attendu_sans_crochets_est_refuse(self):
        self._etapes(
            f"10;saisir;press;wnd[0]/tbar[0]/btn[8];{ECRAN};;;;;S;non;;;")
        with self.assertRaises(TableurInvalide):
            convertir(self.racine)


class TestLeRang(Base):
    """Cliquer sur un en-tete pour trier reordonne les etapes en silence.

    Et l'ordre des etapes est ce qui est tape dans SAP, dans quel ordre, et a
    quel moment on sauvegarde. Un tri sur `nom` produirait une pipeline qui
    charge, qui tourne, et qui fait autre chose.
    """

    def test_un_rang_DECROISSANT_est_refuse(self):
        self._etapes(
            f"20;sauver;press;wnd[0]/tbar[0]/btn[11];{ECRAN};;;;;;oui;;;",
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};colonne;site;;;;non;;;")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("CROITRE", str(capture.exception))

    def test_deux_rangs_EGAUX_sont_refuses(self):
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};colonne;site;;;;non;;;",
            f"10;sauver;press;wnd[0]/tbar[0]/btn[11];{ECRAN};;;;;;oui;;;")
        with self.assertRaises(TableurInvalide):
            convertir(self.racine)

    def test_les_TROUS_sont_permis(self):
        """On numerote de dix en dix pour pouvoir inserer."""
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};colonne;site;;;;non;;;",
            f"90;sauver;press;wnd[0]/tbar[0]/btn[11];{ECRAN};;;;;;oui;;;")
        self.assertIn("sauver", convertir(self.racine))


class TestDerogationOrpheline(Base):
    """Le refus que ce module est SEUL a pouvoir poser.

    Le chargeur ne verifie que le PREFIXE `etape:`. Une faute de frappe donne
    donc une derogation acceptee QUI NE COUVRE RIEN : la garde reste armee,
    l'utilisateur croit l'avoir relachee, et le lot tombe a l'endroit meme ou
    il pensait etre passe.
    """

    def _derogation(self, etape: str, portee: str = "") -> None:
        self._poser(CSV_DEROGATIONS,
                    f"etape;garde;portee;motif\n"
                    f"{etape};relecture;{portee};L'ecran de confirmation "
                    f"n'expose pas le champ relu, releve le 2026-09-08.\n")

    def test_une_derogation_sur_une_etape_INCONNUE_est_refusee(self):
        self._derogation("sauvre")           # faute de frappe
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        message = str(capture.exception)
        self.assertIn("sauvre", message)
        self.assertIn("saisir", message)     # ce qui existe vraiment

    def test_une_PORTEE_qui_designe_une_etape_inconnue_est_refusee(self):
        self._derogation("sauver", portee="etape:sauvre")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("ne couvrirait rien", str(capture.exception))

    def test_une_PORTEE_qui_designe_une_AUTRE_etape_existante_est_refusee(self):
        """Le cas que ce test manquait, et qui est le plus vicieux.

        Les deux etapes existent, donc rien ne signalait la faute. Or le
        controleur n'accorde une derogation que si sa portee vaut `*` ou
        `etape:<son propre nom>` : visant une autre, elle est MORTE PAR
        CONSTRUCTION.

        Et `Contrat.relachements` ecrivait quand meme « derogee » au journal,
        sans consulter la portee — l'humain qui relit voyait une garde
        relachee pendant qu'elle restait armee, et le lot tombait a l'endroit
        meme ou il se croyait passe.
        """
        self._derogation("sauver", portee="etape:saisir")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        message = str(capture.exception)
        self.assertIn("saisir", message)
        self.assertIn("sauver", message)

    def test_la_portee_TOTALE_reste_permise(self):
        """`*` relache la garde sur toute la pipeline : c'est explicite, tres
        visible au recapitulatif, et ca doit rester possible."""
        self._derogation("sauver", portee="*")
        self.assertIn("sauver", convertir(self.racine))

    def test_un_motif_trop_court_est_refuse(self):
        self._poser(CSV_DEROGATIONS,
                    "etape;garde;portee;motif\nsauver;relecture;;trop court\n")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("desactivee en douce", str(capture.exception))

    def test_une_garde_NON_DEROGEABLE_est_refusee(self):
        self._poser(CSV_DEROGATIONS,
                    "etape;garde;portee;motif\n"
                    "sauver;identite;;Motif suffisamment long pour passer le "
                    "seuil minimal impose.\n")
        with self.assertRaises(TableurInvalide):
            convertir(self.racine)


class TestRienNEstEcritSiLeYamlEstRefuse(Base):
    """Rien n'est ecrit avant d'avoir ete recharge. Jamais l'inverse.

    Un YAML casse a cote d'un YAML valide plus ancien, c'est le mauvais
    fichier lance un jour de fatigue.
    """

    def test_aucun_fichier_ne_reste_apres_un_refus(self):
        # `action: sauter` n'existe pas : le CONVERTISSEUR l'accepte — il ne
        # redit pas la liste des actions — et le CHARGEUR le refuse. C'est
        # exactement le partage voulu, et c'est ce qui met la relecture a
        # l'epreuve.
        self._etapes(
            f"10;saisir;sauter;wnd[0]/usr/ctxtW;{ECRAN};;;;;;non;;;")
        sortie = self.racine / "p.yaml"
        with self.assertRaises(TableurInvalide) as capture:
            convertir_fichiers(self.racine, sortie)

        self.assertFalse(sortie.exists(), "un YAML refuse a ete ecrit")
        restes = [c.name for c in self.racine.glob("*.relecture")]
        self.assertEqual(restes, [], "un fichier de relecture a survecu")
        # Le message montre le YAML produit : sans ca, il parle d'un fichier
        # que l'utilisateur du classeur n'a pas ecrit et ne peut pas lire.
        self.assertIn("sauter", str(capture.exception))

    def test_un_YAML_valide_plus_ancien_n_est_PAS_ecrase_par_un_refus(self):
        sortie = convertir_fichiers(self.racine, self.racine / "p.yaml")
        avant = sortie.read_text(encoding="utf-8")

        self._etapes(f"10;saisir;sauter;wnd[0]/usr/ctxtW;{ECRAN};;;;;;non;;;")
        with self.assertRaises(TableurInvalide):
            convertir_fichiers(self.racine, sortie)
        self.assertEqual(sortie.read_text(encoding="utf-8"), avant)


class TestColonnesEtProprietes(Base):

    def test_une_colonne_INCONNUE_est_refusee_et_pas_ignoree(self):
        """« sauvgarde » mal orthographie serait ignore, et l'etape n'ecrirait
        jamais dans SAP : un lot qui semble passer et n'a rien fait."""
        self._poser(CSV_ETAPES, ETAPES.replace("sauvegarde;", "sauvgarde;"))
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("sauvgarde", str(capture.exception))

    def test_une_colonne_SUPPRIMEE_est_refusee(self):
        """La docstring promettait « EXACTEMENT celles attendues » ; le
        controle ne verifiait que l'INCLUSION.

        Mesure avant correction : `sauvegarde` retiree, l'etape `sauver`
        chargeait avec `sauvegarde=False`. Le contrat n'annoncait plus la
        sauvegarde imminente, donc l'etat `douteux` cessait d'exister — et un
        item interrompu apres une sauvegarde repartait a la reprise. Double
        ecriture, par une colonne absente.
        """
        sans = "\n".join(
            ";".join(c for c in ligne.split(";")
                     if not (ligne.startswith("rang") and c == "sauvegarde"))
            for ligne in ETAPES.splitlines())
        # On retire aussi la cellule correspondante de chaque ligne.
        entete = ETAPES.splitlines()[0].split(";")
        rang = entete.index("sauvegarde")
        sans = "\n".join(
            ";".join(c for i, c in enumerate(ligne.split(";")) if i != rang)
            for ligne in ETAPES.splitlines()) + "\n"
        self._poser(CSV_ETAPES, sans)

        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        message = str(capture.exception)
        self.assertIn("sauvegarde", message)
        self.assertIn("absente", message)

    def test_une_colonne_DUPLIQUEE_est_refusee(self):
        """`DictReader` garde la DERNIERE : la valeur de la premiere est
        perdue. Ici la valeur perdue est une CIBLE — la saisie partirait dans
        un autre champ SAP que celui que la feuille montre en tete."""
        entete = ETAPES.splitlines()[0].split(";")
        rang = entete.index("cible")
        double = []
        for ligne in ETAPES.splitlines():
            cellules = ligne.split(";")
            cellules.insert(rang + 1, "cible" if ligne.startswith("rang")
                            else "wnd[0]/usr/ctxtAUTRE")
            double.append(";".join(cellules))
        self._poser(CSV_ETAPES, "\n".join(double) + "\n")

        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("double", str(capture.exception))

    def test_une_propriete_declaree_DEUX_FOIS_est_refusee(self):
        self._poser(CSV_PIPELINE, PIPELINE + "plafond_items;100000\n")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("deux fois", str(capture.exception))

    def test_un_plafond_manquant_est_refuse(self):
        self._poser(CSV_PIPELINE,
                    PIPELINE.replace("plafond_sauvegardes;50\n", ""))
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("plafond_sauvegardes", str(capture.exception))

    def test_un_plafond_non_ENTIER_est_refuse(self):
        self._poser(CSV_PIPELINE,
                    PIPELINE.replace("plafond_items;50", "plafond_items;tous"))
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("rayon d'action", str(capture.exception))

    def test_une_feuille_d_etapes_VIDE_est_refusee(self):
        self._poser(CSV_ETAPES, ENTETE)
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("se charge et ne fait rien", str(capture.exception))

    def test_sauvegarde_attend_oui_ou_non_pas_VRAI_FAUX(self):
        """Le VRAI/FAUX d'Excel est localise et ne se relit pas d'une machine
        a l'autre."""
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;{ECRAN};colonne;site;;;;VRAI;;;")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("localise", str(capture.exception))

    def test_un_ecran_MALFORME_est_refuse(self):
        self._etapes(
            "10;saisir;set;wnd[0]/usr/ctxtW;IA08/RIPLKO10/1000;colonne;site;"
            ";;;non;;;")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("jeton d'ecran", str(capture.exception))

    def test_un_ecran_VIDE_est_refuse(self):
        self._etapes(
            f"10;saisir;set;wnd[0]/usr/ctxtW;;colonne;site;;;;non;;;")
        with self.assertRaises(TableurInvalide) as capture:
            convertir(self.racine)
        self.assertIn("garde d'identite", str(capture.exception))


class TestLExempleLivre(unittest.TestCase):
    """`exemple/` doit rester convertible, chargeable et executable.

    C'est de la documentation qui TOURNE. Un exemple que rien ne verifie se
    perime — et celui-la est le premier fichier que quelqu'un copie pour
    ecrire son propre automatisme : s'il derive, il enseigne la derive.
    """

    RACINE = Path(__file__).resolve().parent.parent / "exemple"

    def test_les_CSV_livres_donnent_une_pipeline_qui_charge(self):
        with tempfile.TemporaryDirectory() as dossier:
            chemin = convertir_fichiers(self.RACINE,
                                        Path(dossier) / "exemple.yaml")
            pipeline = charger(chemin)
        self.assertEqual(pipeline.nom, "exemple_variantes")
        self.assertEqual(len(pipeline.etapes), 5)
        self.assertTrue(any(e.sauvegarde for e in pipeline.etapes),
                        "une pipeline qui ne sauvegarde nulle part n'ecrit "
                        "rien dans SAP")

    def test_le_jeu_livre_se_lit_avec_les_cles_de_la_pipeline(self):
        from falcon.donnees import lire_items
        with tempfile.TemporaryDirectory() as dossier:
            pipeline = charger(convertir_fichiers(
                self.RACINE, Path(dossier) / "exemple.yaml"))
        items, _ = lire_items(self.RACINE / "jeu.csv", pipeline.cles)
        self.assertEqual(len(items), 3)

    def test_le_script_de_demonstration_va_AU_BOUT(self):
        """Il ne suffit pas qu'il existe : il doit tourner, ici, sans SAP."""
        import subprocess
        rendu = subprocess.run(
            [sys.executable, str(self.RACINE / "rejouer.py")],
            capture_output=True, text=True, timeout=120)
        self.assertEqual(rendu.returncode, 0, rendu.stderr)
        # Ce que la demonstration promet de montrer.
        for attendu in ("000000000010023456",     # zeros: 18 apres elagage
                        "/BCP01_K75",             # gabarit + majuscules
                        "ok': 3"):                # les trois items passent
            with self.subTest(attendu=attendu):
                self.assertIn(attendu, rendu.stdout)


class TestLeClasseurLivre(unittest.TestCase):
    """La couture entre le classeur et le convertisseur.

    C'est l'endroit qui derive : les colonnes vivent dans `falcon/tableur/`,
    le classeur les recopie, et rien ne les faisait s'accorder. Un en-tete qui
    prend un « s » d'un cote produit un CSV dont la colonne est INCONNUE — ou
    pire, une colonne attendue et absente.

    Le classeur est un binaire livre dans le depot : ces tests le lisent tel
    qu'il est, pas tel qu'on le reconstruirait.
    """

    CLASSEUR = (Path(__file__).resolve().parent.parent
                / "classeur" / "FALCON.xlsx")

    def setUp(self):
        if not self.CLASSEUR.exists():
            self.skipTest("classeur/FALCON.xlsx absent")
        try:
            import openpyxl                              # noqa: F401
        except ImportError:
            # `openpyxl` est un outil de CONSTRUCTION, pas une dependance
            # d'execution : la regle « PyYAML seule dependance » tient, et un
            # poste qui ne l'a pas doit pouvoir lancer la suite. `verifier.py`
            # liste ce qui n'a pas tourne.
            self.skipTest("openpyxl absent : outil de construction, pas "
                          "dependance d'execution")

    def test_le_classeur_est_un_zip_dont_tout_le_XML_se_parse(self):
        """Ce que je PEUX affirmer sans ouvrir Excel. Excel refuse le fichier
        entier si une seule partie est mal formee."""
        import xml.etree.ElementTree as ET
        import zipfile

        archive = zipfile.ZipFile(self.CLASSEUR)
        self.assertIsNone(archive.testzip())
        parties = [n for n in archive.namelist()
                   if n.endswith((".xml", ".rels"))]
        self.assertGreater(len(parties), 5, "archive suspecte")
        for nom in parties:
            with self.subTest(partie=nom):
                ET.fromstring(archive.read(nom))

    def _entete(self, feuille: str, ligne: int = 2) -> tuple[str, ...]:
        from openpyxl import load_workbook
        classeur = load_workbook(self.CLASSEUR, read_only=True)
        valeurs = [c.value for c in classeur[feuille][ligne]]
        classeur.close()
        return tuple(v for v in valeurs if v is not None)

    def test_les_en_tetes_sont_EXACTEMENT_ceux_du_convertisseur(self):
        for feuille, attendues in (("Etapes", COLONNES_ETAPES),
                                   ("Derogations", COLONNES_DEROGATIONS)):
            with self.subTest(feuille=feuille):
                self.assertEqual(self._entete(feuille), attendues)

    def test_la_feuille_EXEMPLE_n_est_pas_exportee(self):
        """Une ligne d'exemple laissee dans une feuille exportee deviendrait
        une ETAPE, et la pipeline porterait une saisie que personne n'a
        voulue."""
        from openpyxl import load_workbook
        classeur = load_workbook(self.CLASSEUR, read_only=True)
        self.assertIn("Exemple", classeur.sheetnames)
        classeur.close()

        macro = (self.CLASSEUR.parent / "ExportFalcon.bas").read_text(
            encoding="utf-8")
        for feuille in ("Pipeline", "Etapes", "Derogations", "Donnees"):
            self.assertIn(f'"{feuille}"', macro, f"{feuille} n'est pas exportee")
        self.assertNotIn('"Exemple"', macro)
        self.assertNotIn('"Dictionnaire"', macro)

    def test_la_feuille_Etapes_est_VIDE_sous_son_en_tete(self):
        """Meme raison : ce qui est sous l'en-tete part dans le CSV."""
        from openpyxl import load_workbook
        classeur = load_workbook(self.CLASSEUR, read_only=True)
        premiere = [c.value for c in classeur["Etapes"][3]]
        classeur.close()
        self.assertTrue(all(v is None for v in premiere),
                        f"la premiere ligne d'Etapes n'est pas vide : "
                        f"{premiere}")

    def test_l_EXEMPLE_du_classeur_produit_une_pipeline_QUI_CHARGE(self):
        """La preuve la plus forte possible sans ouvrir Excel.

        On lit le classeur, on ecrit les CSV comme la macro les ecrirait —
        meme en-tete en ligne 2, memes colonnes, lignes vides et lignes « # »
        sautees — et on convertit. Si le classeur et le convertisseur ne
        s'accordent plus, ca tombe ici, avec le message du convertisseur.
        """
        from openpyxl import load_workbook

        classeur = load_workbook(self.CLASSEUR)

        def exporter(feuille: str, vers: Path, entete: int = 2) -> None:
            lignes = []
            for rang in classeur[feuille].iter_rows(min_row=entete,
                                                    values_only=True):
                cellules = ["" if v is None else str(v) for v in rang]
                if not any(c.strip() for c in cellules):
                    continue
                if cellules[0].startswith("#"):
                    continue
                lignes.append(";".join(cellules))
            vers.write_text("\n".join(lignes) + "\n", encoding="utf-8")

        with tempfile.TemporaryDirectory() as dossier:
            racine = Path(dossier)
            exporter("Pipeline", racine / CSV_PIPELINE)
            # C'est la feuille EXEMPLE qui porte des etapes : `Etapes` est
            # livree vide, justement pour ne rien exporter qu'on n'ait voulu.
            exporter("Exemple", racine / CSV_ETAPES)

            # Le classeur livre `Pipeline` sans valeurs — c'est a l'
            # utilisateur de les mettre. On les fournit ici, pour que le test
            # porte sur les ETAPES, qui sont ce que le classeur propose.
            (racine / CSV_PIPELINE).write_text(
                "propriete;valeur\nnom;depuis_le_classeur\n"
                "classe;iterative\ncles;site\nplafond_items;50\n"
                "plafond_sauvegardes;50\n", encoding="utf-8")

            pipeline = charger(convertir_fichiers(racine,
                                                  racine / "p.yaml"))

        self.assertEqual([e.nom for e in pipeline.etapes],
                         ["saisir_site", "saisir_variante",
                          "saisir_equipement", "sauver"])
        self.assertTrue(pipeline.etapes[-1].sauvegarde)

    def test_la_macro_ne_VALIDE_rien(self):
        """La parade au VBA non teste est architecturale : le garder bete.

        Une macro fausse doit produire un CSV REFUSE avec un message situe,
        pas un YAML plausible. Si elle validait, une erreur dedans deviendrait
        une pipeline qui se charge et fait autre chose.
        """
        macro = (self.CLASSEUR.parent / "ExportFalcon.bas").read_text(
            encoding="utf-8")
        # Elle n'a aucune raison de connaitre le vocabulaire du domaine.
        for mot in ("iterative", "navigation_libre", "connue_benigne",
                    "plafond_items"):
            with self.subTest(mot=mot):
                self.assertNotIn(mot, macro)


if __name__ == "__main__":
    unittest.main()
